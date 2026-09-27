#!/usr/bin/env python3
"""Read local sources without changing them; regenerate the provenance manifest.

Install requirements-ingest.txt in an isolated environment before running.
Extracted text is a reading aid, never a replacement for original layout.
"""
from pathlib import Path
from collections import defaultdict
import hashlib
import json
import struct
import zlib

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "wiki/sources/extracted"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_id(path, digest):
    if path.suffix == ".pdf":
        if path.stem.startswith("TTAK"):
            return "tta-r3"
        if path.stem.isdigit():
            return "lecture-" + path.stem
        return "pdf-" + digest[:12]
    return path.suffix[1:] + "-" + digest[:12]


def extract_xls(path):
    import olefile
    import xlrd
    from xlrd.formula import decompile_formula, FMLA_TYPE_CELL, colname
    book = xlrd.open_workbook(path)
    with olefile.OleFileIO(path) as ole:
        stream = ole.openstream("Workbook" if ole.exists("Workbook") else "Book").read()
    offsets = []
    pos = 0
    while pos + 4 <= len(stream):
        code, size = struct.unpack_from("<HH", stream, pos)
        body = stream[pos + 4:pos + 4 + size]
        if code == 0x85:
            offsets.append(struct.unpack_from("<I", body)[0])
        pos += 4 + size
        if code == 0x0A:
            break
    sheets = []
    for sheet, offset in zip(book.sheets(), offsets, strict=True):
        formulas = {}
        pos = offset
        while pos + 4 <= len(stream):
            code, size = struct.unpack_from("<HH", stream, pos)
            body = stream[pos + 4:pos + 4 + size]
            if code == 0x06 and len(body) >= 22:
                row, col = struct.unpack_from("<HH", body)
                length = struct.unpack_from("<H", body, 20)[0]
                try:
                    formula = decompile_formula(book, body[22:22 + length], length,
                                                FMLA_TYPE_CELL, row, col)
                    formulas[(row, col)] = {"formula": formula, "formula_status": "decompiled-not-recalculated"}
                except Exception as exc:
                    formulas[(row, col)] = {"formula_status": "unresolved", "error": type(exc).__name__}
            pos += 4 + size
            if code == 0x0A:
                break
        cells = []
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                value = sheet.cell_value(row, col)
                if value != "" or (row, col) in formulas:
                    cells.append({"cell": f"{colname(col)}{row + 1}", "cached_value": value,
                                  **formulas.get((row, col), {})})
        sheets.append({"name": sheet.name, "rows": sheet.nrows, "columns": sheet.ncols, "cells": cells})
    return {"sheets": sheets, "limitation": "Cached values and BIFF formula text only; no recalculation, macros, controls, or visual verification."}


def extract_xlsx(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in book:
        cells = []
        used_rows = used_columns = 0
        for row in sheet:
            for cell in row:
                if cell.value is None:
                    continue
                used_rows = max(used_rows, cell.row)
                used_columns = max(used_columns, cell.column)
                item = {"cell": cell.coordinate, "value": cell.value}
                if cell.data_type == "f":
                    item["formula"] = item.pop("value")
                    item["cached_value"] = cached[sheet.title][cell.coordinate].value
                cells.append(item)
        sheets.append({"name": sheet.title, "rows": sheet.max_row or used_rows,
                       "columns": sheet.max_column or used_columns, "cells": cells})
    book.close()
    cached.close()
    return {"sheets": sheets, "limitation": "Cell content only; drawings, formatting and native recalculation not verified."}


def extract_hwp(path):
    import olefile
    with olefile.OleFileIO(path) as ole:
        preview = ole.openstream("PrvText").read().decode("utf-16le", errors="replace") if ole.exists("PrvText") else ""
        header = ole.openstream("FileHeader").read()
        flags = struct.unpack_from("<I", header, 36)[0]
        if flags & 2:
            return {"preview": preview, "sections": [], "limitation": "Encrypted body; preview only."}
        sections = []
        for stream in sorted(ole.listdir()):
            if len(stream) != 2 or stream[0] != "BodyText":
                continue
            data = ole.openstream(stream).read()
            if flags & 1:
                data = zlib.decompress(data, -15)
            paragraphs = []
            pos = 0
            while pos + 4 <= len(data):
                record = struct.unpack_from("<I", data, pos)[0]
                pos += 4
                tag, size = record & 0x3FF, record >> 20
                if size == 0xFFF:
                    size = struct.unpack_from("<I", data, pos)[0]
                    pos += 4
                body = data[pos:pos + size]
                pos += size
                if tag != 67:  # HWPTAG_PARA_TEXT
                    continue
                units = struct.unpack("<" + "H" * (len(body) // 2), body[:len(body) // 2 * 2])
                plain = []
                i = 0
                while i < len(units):
                    u = units[i]
                    if u < 32:
                        if u in (10, 13):
                            plain.append("\n")
                        elif u == 9:
                            plain.append("\t")
                        i += 8 if u in (1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23) else 1
                    else:
                        plain.append(chr(u))
                        i += 1
                text = "".join(plain).strip()
                if text:
                    paragraphs.append(text)
            sections.append({"name": "/".join(stream), "paragraphs": paragraphs})
    return {"preview": preview, "sections": sections,
            "limitation": "Best-effort HWP v5 paragraph extraction; table relationships, embedded objects, pagination and visual layout are not verified. Preview may be truncated."}


def main():
    from pypdf import PdfReader
    OUT.mkdir(parents=True, exist_ok=True)
    groups = defaultdict(list)
    for path in sorted((ROOT / "reference").rglob("*")):
        if path.is_file() and not path.name.startswith("."):
            groups[sha(path)].append(path)
    records = []
    for digest, paths in groups.items():
        path = min(paths, key=lambda p: (len(p.parts), str(p)))
        sid = source_id(path, digest)
        if path.suffix == ".pdf":
            reader = PdfReader(path)
            content = {"pages": [{"page": i, "text": p.extract_text(extraction_mode="layout") or ""}
                                 for i, p in enumerate(reader.pages, 1)],
                       "limitation": "Text layer only; images and layout must be checked against the original PDF."}
        elif path.suffix == ".xlsx":
            content = extract_xlsx(path)
        elif path.suffix == ".xls":
            content = extract_xls(path)
        elif path.suffix == ".hwp":
            content = extract_hwp(path)
        else:
            raise ValueError(f"Unsupported source: {path}")
        target = OUT / f"{sid}.json"
        target.write_text(json.dumps({"source_id": sid, **content}, ensure_ascii=False, indent=2, default=str) + "\n")
        record = {"id": sid, "path": str(path.relative_to(ROOT)), "sha256": digest,
                  "aliases": [str(p.relative_to(ROOT)) for p in paths if p != path],
                  "bytes": path.stat().st_size, "format": path.suffix[1:],
                  "extraction": str(target.relative_to(ROOT)), "extraction_sha256": sha(target),
                  "card": f"wiki/sources/{sid}.md", "limitation": content["limitation"]}
        if "pages" in content:
            record["page_count"] = len(content["pages"])
        if "sheets" in content:
            record["sheets"] = [{"name": s["name"], "rows": s["rows"], "columns": s["columns"]} for s in content["sheets"]]
        records.append(record)
    manifest = {"schema_version": 1, "inventory_date": "2026-09-27", "scope": "local-reference",
                "original_file_count": sum(len(r["aliases"]) + 1 for r in records),
                "unique_source_count": len(records), "sources": sorted(records, key=lambda r: r["id"])}
    (ROOT / "wiki/sources/manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(f"Ingested {manifest['original_file_count']} files / {len(records)} unique sources; originals unchanged.")


if __name__ == "__main__":
    main()
